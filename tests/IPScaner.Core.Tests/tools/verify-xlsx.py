"""Out-of-process validator for the workbook written by TableExporter.WriteXlsx.

Usage:
    <bundled python> verify-xlsx.py <path to .xlsx>

Prints every sheet name and cell so the values written by the C# exporter can be
compared with what a real spreadsheet library reads back.
"""

import sys

import openpyxl


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: verify-xlsx.py <workbook.xlsx>")
        return 2

    path = sys.argv[1]
    print(f"openpyxl {openpyxl.__version__} reading {path}")

    workbook = openpyxl.load_workbook(path)
    print(f"sheetnames: {workbook.sheetnames}")

    for name in workbook.sheetnames:
        sheet = workbook[name]
        print(f"--- sheet {name!r} dimension={sheet.calculate_dimension()} max_row={sheet.max_row} max_column={sheet.max_column}")
        for row in sheet.iter_rows():
            for cell in row:
                print(f"  {cell.coordinate}: value={cell.value!r} type={cell.data_type}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
